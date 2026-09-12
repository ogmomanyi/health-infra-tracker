"""Evidence identity, document discovery, and deterministic text extraction."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import csv
import json
import re
from pathlib import Path
from typing import Any, Mapping
from urllib.parse import urljoin, urlparse
import zipfile
import xml.etree.ElementTree as ET

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader


MATERIAL_FIELDS = (
    "source",
    "source_record_id",
    "tender_reference",
    "title",
    "buyer",
    "country",
    "publication_date",
    "closing_date",
    "project_reference",
    "procurement_stage",
    "notice_text",
    "document_urls",
)

DOCUMENT_SUFFIXES = {
    ".csv",
    ".doc",
    ".docx",
    ".htm",
    ".html",
    ".pdf",
    ".txt",
    ".xls",
    ".xlsx",
}

DOCUMENT_LINK_TERMS = (
    "attachment",
    "bid document",
    "bidding document",
    "download",
    "procurement document",
    "request for bids",
    "request for quotation",
    "specification",
    "terms of reference",
    "tender document",
)


def is_valid_document_url(url: object) -> bool:
    """Reject known listing/navigation pages that cannot evidence one notice."""
    value = normalized_text(url)
    if not value:
        return False
    parsed = urlparse(value)
    host = parsed.netloc.casefold().split(":", 1)[0]
    path = parsed.path.casefold().rstrip("/")
    if host.endswith("procurement-notices.undp.org") and path in {
        "/index.cfm", "/search.cfm",
    }:
        return False
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def normalized_text(value: object) -> str:
    return " ".join(str(value or "").split())


def _identity_token(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", normalized_text(value).casefold()).strip()


def stable_process_id(record: Mapping[str, Any]) -> str:
    """Identify a contracting process without depending on a mutable title."""
    source = _identity_token(record.get("source"))
    identity = next(
        (
            _identity_token(record.get(field))
            for field in ("source_record_id", "tender_reference", "source_url")
            if _identity_token(record.get(field))
        ),
        _identity_token(record.get("title")),
    )
    key = f"{source}|{identity}"
    return "proc_process_" + sha256(key.encode("utf-8")).hexdigest()[:20]


def evidence_content_hash(record: Mapping[str, Any]) -> str:
    material = {
        field: normalized_text(record.get(field))
        for field in MATERIAL_FIELDS
    }
    payload = json.dumps(material, ensure_ascii=True, sort_keys=True, separators=(",", ":"))
    return sha256(payload.encode("utf-8")).hexdigest()


def stable_release_id(process_id: str, content_hash: str) -> str:
    key = f"{process_id}|{content_hash}"
    return "proc_release_" + sha256(key.encode("utf-8")).hexdigest()[:20]


def enrich_evidence_identity(record: Mapping[str, Any]) -> dict[str, Any]:
    enriched = dict(record)
    process_id = normalized_text(enriched.get("procurement_process_id")) or stable_process_id(enriched)
    content_hash = normalized_text(enriched.get("content_hash")) or evidence_content_hash(enriched)
    enriched["procurement_process_id"] = process_id
    enriched["content_hash"] = content_hash
    enriched["procurement_release_id"] = (
        normalized_text(enriched.get("procurement_release_id"))
        or stable_release_id(process_id, content_hash)
    )
    return enriched


def discover_document_links(html: str, base_url: str) -> list[str]:
    """Return unique document-like links from an official notice page."""
    soup = BeautifulSoup(html or "", "html.parser")
    links: list[str] = []
    seen: set[str] = set()
    for anchor in soup.select("a[href]"):
        href = normalized_text(anchor.get("href"))
        if not href or href.startswith(("#", "javascript:", "mailto:")):
            continue
        url = urljoin(base_url, href)
        if not is_valid_document_url(url):
            continue
        path = Path(urlparse(url).path)
        label = normalized_text(anchor.get_text(" ", strip=True)).casefold()
        if path.suffix.casefold() not in DOCUMENT_SUFFIXES and not any(
            term in label for term in DOCUMENT_LINK_TERMS
        ):
            continue
        if url not in seen:
            seen.add(url)
            links.append(url)
    return links


def encode_document_urls(urls: list[str]) -> str:
    return json.dumps(list(dict.fromkeys(urls)), ensure_ascii=True)


def decode_document_urls(value: object) -> list[str]:
    text = normalized_text(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        parsed = re.split(r"[;|\n]+", text)
    if isinstance(parsed, str):
        parsed = [parsed]
    if not isinstance(parsed, list):
        return []
    return list(dict.fromkeys(
        normalized_text(item)
        for item in parsed
        if is_valid_document_url(item)
    ))


def html_text(payload: bytes | str) -> str:
    if isinstance(payload, bytes):
        payload = payload.decode("utf-8", errors="replace")
    soup = BeautifulSoup(payload, "html.parser")
    for node in soup.select("script, style, nav, footer, noscript"):
        node.decompose()
    return normalized_text(soup.get_text(" ", strip=True))


def _docx_text(payload: bytes) -> str:
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        document = archive.read("word/document.xml")
    root = ET.fromstring(document)
    paragraphs = [
        normalized_text(" ".join(node.itertext()))
        for node in root.iter()
        if node.tag.rsplit("}", 1)[-1] == "p"
    ]
    return "\n".join(paragraph for paragraph in paragraphs if paragraph)


def _xlsx_text(payload: bytes) -> str:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [normalized_text(value) for value in row]
                if any(values):
                    lines.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(lines)


def _csv_text(payload: bytes) -> str:
    decoded = payload.decode("utf-8-sig", errors="replace")
    rows = csv.reader(decoded.splitlines())
    return "\n".join(" | ".join(normalized_text(value) for value in row) for row in rows)


def extract_document_text(
    payload: bytes,
    *,
    content_type: str = "",
    source_url: str = "",
) -> tuple[str, str]:
    """Return extracted text and a provenance-friendly extraction status."""
    suffix = Path(urlparse(source_url).path).suffix.casefold()
    mime = content_type.split(";", 1)[0].strip().casefold()
    try:
        if suffix == ".pdf" or mime == "application/pdf":
            reader = PdfReader(BytesIO(payload))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            text = normalized_text(text)
            return (text, "EXTRACTED") if text else ("", "OCR_REQUIRED")
        if suffix == ".docx" or mime.endswith("wordprocessingml.document"):
            return normalized_text(_docx_text(payload)), "EXTRACTED"
        if suffix in {".xlsx", ".xlsm"} or "spreadsheetml" in mime:
            return normalized_text(_xlsx_text(payload)), "EXTRACTED"
        if suffix == ".csv" or mime in {"text/csv", "application/csv"}:
            return normalized_text(_csv_text(payload)), "EXTRACTED"
        if suffix in {".html", ".htm"} or mime in {"text/html", "application/xhtml+xml"}:
            return html_text(payload), "EXTRACTED"
        if mime.startswith("text/") or suffix == ".txt":
            return normalized_text(payload.decode("utf-8", errors="replace")), "EXTRACTED"
        return "", "UNSUPPORTED"
    except (KeyError, OSError, ValueError, zipfile.BadZipFile):
        return "", "EXTRACTION_FAILED"
